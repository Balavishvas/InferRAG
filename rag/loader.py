from pathlib import Path

import fitz

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx", ".csv",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp",
}


def _load_pdf(path: Path) -> str:
    with fitz.open(path) as doc:
        pages = []
        for page_number, page in enumerate(doc, 1):
            text = page.get_text().strip()
            if text:
                pages.append(f"[Page {page_number}]\n{text}")
        return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_number, table in enumerate(document.tables, 1):
        parts.append(f"[Table {table_number}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def _load_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    slides = []

    for slide_number, slide in enumerate(presentation.slides, 1):
        parts = [f"[Slide {slide_number}]"]

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())

            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    parts.append(" | ".join(cells))

        slides.append("\n".join(parts))

    return "\n\n".join(slides)


def _load_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []

    try:
        for worksheet in workbook.worksheets:
            rows = [f"[Sheet: {worksheet.title}]"]
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(" | ".join(values))
            sheets.append("\n".join(rows))
    finally:
        workbook.close()

    return "\n\n".join(sheets)


def _load_image(path: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise RuntimeError(
            "Image ingestion requires Pillow and pytesseract. "
            "Install them with: pip install Pillow pytesseract"
        ) from exc

    try:
        image = Image.open(path)
        text = pytesseract.image_to_string(image).strip()
    except Exception as exc:
        raise RuntimeError(
            "Image OCR failed. Make sure Tesseract OCR is installed and available "
            "on PATH."
        ) from exc

    if not text:
        raise ValueError(
            f"No readable text was found in image: {path.name}. "
            "Vision-based image understanding will be added in a later multimodal stage."
        )

    return f"[Image: {path.name}]\n{text}"


def load_document(path: str) -> str:
    """Load supported files into normalized text for the RAG pipeline."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported file type: {suffix or '[no extension]'}. "
            f"Supported types: {supported}"
        )

    if suffix in {".txt", ".md", ".csv"}:
        return file_path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix == ".docx":
        return _load_docx(file_path)
    if suffix == ".pptx":
        return _load_pptx(file_path)
    if suffix == ".xlsx":
        return _load_xlsx(file_path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        return _load_image(file_path)

    raise ValueError(f"Unsupported file type: {suffix}")
